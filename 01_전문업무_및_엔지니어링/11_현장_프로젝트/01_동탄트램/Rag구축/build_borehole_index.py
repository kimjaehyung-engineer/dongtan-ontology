# -*- coding: utf-8 -*-
import os, sys, json
import openpyxl

# 1. Borehole mappings for 1공구 (54 pages, 45 holes)
nh_gb_pages = [
    # (공번, 시작페이지, 끝페이지, 공구/위치, 기본표고)
    ('NH-1', 1, 2, '1공구 본선', 21.37),
    ('NH-2', 3, 3, '1공구 본선', 20.53),
    ('NH-3', 4, 4, '1공구 본선', None),
    ('NH-4', 5, 5, '1공구 본선', 26.96),
    ('NH-5', 6, 6, '1공구 본선', 29.50),
    ('NH-6', 7, 7, '1공구 본선', 31.30),
    ('NH-7', 8, 8, '1공구 본선', 36.50),
    ('NH-8', 9, 9, '1공구 본선', 54.26),
    ('NH-9', 10, 10, '1공구 본선', 49.27),
    ('NH-10', 11, 11, '1공구 본선', 46.25),
    ('NH-11', 12, 12, '1공구 본선', None),
    ('NH-12', 13, 13, '1공구 본선', 42.70),
    ('NH-13', 14, 14, '1공구 본선', 35.07),
    ('NH-14', 15, 16, '1공구 본선', 45.92),
    ('NH-15', 17, 17, '1공구 본선', 52.44),
    ('NH-16', 18, 18, '1공구 본선', 48.19),
    ('NH-17', 19, 19, '1공구 본선', None),
    ('NH-18', 20, 20, '1공구 본선', 52.86),
    ('NH-19', 21, 21, '1공구 본선', 64.48),
    ('NH-20', 22, 22, '1공구 본선', 62.38),
    ('NH-21', 23, 23, '1공구 본선', None),
    ('NH-22', 24, 24, '1공구 본선', None),
    ('NH-23', 25, 25, '1공구 본선', 90.56),
    ('NH-24', 26, 27, '1공구 본선', 78.48),
    ('NH-25', 28, 28, '1공구 본선', 59.23),
    ('NH-26', 29, 29, '1공구 본선', 55.66),
    ('NH-27', 30, 30, '1공구 본선', 77.22),
    ('NH-28', 31, 31, '1공구 본선', 63.45),
    ('NH-29', 32, 32, '1공구 본선', 76.68),
    ('NH-30', 33, 33, '1공구 본선', 63.17),
    ('NH-31', 34, 34, '1공구 본선', None),
    ('NH-32', 35, 35, '1공구 본선', None),
    ('NH-33', 36, 36, '1공구 본선', None),
    ('NH-34', 37, 37, '1공구 본선', None),
    ('NH-35', 38, 38, '1공구 본선', None),
    ('NH-36', 39, 39, '1공구 본선', None),
    ('GB-1', 40, 41, '차량기지', 49.05),
    ('GB-2', 42, 42, '차량기지', 50.60),
    ('GB-3', 43, 43, '차량기지', 50.63),
    ('GB-4', 44, 45, '차량기지', 57.63),
    ('GB-5', 46, 47, '차량기지', 56.45),
    ('GB-6', 48, 49, '차량기지', 54.83),
    ('GB-7', 50, 51, '차량기지', 54.74),
    ('GB-8', 52, 52, '차량기지', 61.45),
    ('GB-9', 53, 54, '차량기지', 68.61),
]

# 2. Borehole mappings for 2공구 (30 pages, 27 holes)
dt_pages = [
    ('DT-1', 1, 2, '2공구 본선', 32.82),
    ('DT-2', 3, 3, '2공구 본선', 38.87),
    ('DT-3', 4, 4, '2공구 본선', 42.78),
    ('DT-4', 5, 5, '2공구 본선', 41.60),
    ('DT-5', 6, 6, '2공구 본선', 47.96),
    ('DT-6', 7, 7, '2공구 본선', 61.69),
    ('DT-7', 8, 8, '2공구 본선', 53.58),
    ('DT-8', 9, 10, '2공구 본선', 39.17),
    ('DT-9', 11, 12, '2공구 본선', 42.40),
    ('DT-10', 13, 13, '2공구 본선', 38.51),
    ('DT-11', 14, 14, '2공구 본선', 48.73),
    ('DT-12', 15, 15, '2공구 본선', 48.81),
    ('DT-13', 16, 16, '2공구 본선', 45.35),
    ('DT-14', 17, 17, '2공구 본선', 36.47),
    ('DT-15', 18, 18, '2공구 본선', 34.38),
    ('DT-16', 19, 19, '2공구 본선', 45.34),
    ('DT-17', 20, 20, '2공구 본선', 61.80),
    ('DT-18', 21, 21, '2공구 본선', 64.35),
    ('DT-19', 22, 22, '2공구 본선', 52.79),
    ('DT-20', 23, 23, '2공구 본선', 50.12),
    ('DT-21', 24, 24, '2공구 본선', 44.10),
    ('DT-22', 25, 25, '2공구 본선', 25.64),
    ('DT-23', 26, 26, '2공구 본선', 24.24),
    ('DT-24', 27, 27, '2공구 본선', 23.78),
    ('DT-25', 28, 28, '2공구 본선', 20.39),
    ('DT-26', 29, 29, '2공구 본선', 19.06),
    ('DT-27', 30, 30, '2공구 본선', 17.04),
]

# Build JSON Index Structure
borehole_index_1 = {
    "document_name": "기본설계 시추주상도(1공구)_45공.pdf",
    "total_pages": 54,
    "total_boreholes": 45,
    "description": "동탄 도시철도(트램) 기본설계 1공구 본선(NH-1~36, 36공) 및 차량기지 부지(GB-1~9, 9공) 총 45개 시추주상도",
    "sections": []
}

for bh, s, e, loc, el in nh_gb_pages:
    borehole_index_1["sections"].append({
        "section_id": bh,
        "title": f"{bh} 시추주상도 ({loc})",
        "facility": [loc, bh],
        "keywords": [bh, "시추주상도", "1공구", loc, "N치", "지반표고", "지하수위", "지층"],
        "start_page": s,
        "end_page": e,
        "elevation_m": el,
        "content_type": "BOREHOLE_LOG"
    })

borehole_index_2 = {
    "document_name": "기본설계 시추주상도(2공구)_27공.pdf",
    "total_pages": 30,
    "total_boreholes": 27,
    "description": "동탄 도시철도(트램) 기본설계 2공구 본선(DT-1~27, 총 27공) 시추주상도 (CAD 벡터 선화)",
    "sections": []
}

for bh, s, e, loc, el in dt_pages:
    borehole_index_2["sections"].append({
        "section_id": bh,
        "title": f"{bh} 시추주상도 ({loc})",
        "facility": [loc, bh],
        "keywords": [bh, "시추주상도", "2공구", "DT", "N치", "지반표고", "지하수위", "지층"],
        "start_page": s,
        "end_page": e,
        "elevation_m": el,
        "content_type": "BOREHOLE_LOG"
    })

# 1. Save JSON index to 04_METADATA
with open('04_METADATA/doc_index_boreholes_1.json', 'w', encoding='utf-8') as f:
    json.dump(borehole_index_1, f, ensure_ascii=False, indent=2)

with open('04_METADATA/doc_index_boreholes_2.json', 'w', encoding='utf-8') as f:
    json.dump(borehole_index_2, f, ensure_ascii=False, indent=2)

# Combined index for easy reference
combined_index = {
    "1공구_및_차량기지": borehole_index_1,
    "2공구": borehole_index_2
}
with open('04_METADATA/doc_index_boreholes.json', 'w', encoding='utf-8') as f:
    json.dump(combined_index, f, ensure_ascii=False, indent=2)

print("Saved JSON indexes to 04_METADATA successfully.")

# 2. Register to server documents _metadata_index.json
server_meta_file = r'C:\Users\sskjh\antigravity\01_전문업무_및_엔지니어링\14_개발_및_자동화_앱\RAG\documents\_metadata_index.json'
server_meta = {}
if os.path.exists(server_meta_file):
    try:
        with open(server_meta_file, 'r', encoding='utf-8') as f:
            server_meta = json.load(f)
    except:
        server_meta = {}

server_meta[borehole_index_1["document_name"]] = borehole_index_1
server_meta[borehole_index_2["document_name"]] = borehole_index_2

with open(server_meta_file, 'w', encoding='utf-8') as f:
    json.dump(server_meta, f, ensure_ascii=False, indent=2)

print("Registered to server _metadata_index.json successfully.")

# 3. Save Markdown Table Guide
md_lines = [
    "# 동탄 도시철도(트램) 기본설계 시추주상도 전수 색인표 (총 72공)",
    "\n본 색인은 1공구(NH 36공), 차량기지(GB 9공), 2공구(DT 27공) 총 72개 시추공의 주상도 수록 페이지 매핑 가이드입니다.\n",
    "## 📌 1공구 및 차량기지 (기본설계 시추주상도(1공구)_45공.pdf - 54p)",
    "| 시추공번 | 대상구간 | 페이지 범위 | 지반표고(EL) | 비고 |",
    "| :--- | :--- | :---: | :---: | :--- |"
]
for item in borehole_index_1["sections"]:
    el_str = f"{item['elevation_m']} m" if item['elevation_m'] is not None else "-"
    md_lines.append(f"| **{item['section_id']}** | {item['facility'][0]} | **p.{item['start_page']} ~ p.{item['end_page']}** | {el_str} | 주상도 |")

md_lines.extend([
    "\n## 📌 2공구 본선 (기본설계 시추주상도(2공구)_27공.pdf - 30p)",
    "| 시추공번 | 대상구간 | 페이지 범위 | 지반표고(EL) | 비고 |",
    "| :--- | :--- | :---: | :---: | :--- |"
])
for item in borehole_index_2["sections"]:
    el_str = f"{item['elevation_m']} m" if item['elevation_m'] is not None else "-"
    md_lines.append(f"| **{item['section_id']}** | {item['facility'][0]} | **p.{item['start_page']} ~ p.{item['end_page']}** | {el_str} | 주상도 |")

with open('04_METADATA/doc_index_boreholes.md', 'w', encoding='utf-8') as f:
    f.write('\n'.join(md_lines) + '\n')

print("Saved Markdown guide to 04_METADATA/doc_index_boreholes.md successfully.")

# 4. Populate borehole_register.xlsx
excel_path = r'04_METADATA\borehole_register.xlsx'
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

# Clear existing data rows (keep header)
while ws.max_row > 1:
    ws.delete_rows(2)

# Insert 1공구 & 차량기지
for item in borehole_index_1["sections"]:
    row = [
        item['section_id'],               # BH_No
        '기본설계',                        # Stage
        '',                               # Coordinate_X
        '',                               # Coordinate_Y
        item['elevation_m'],              # Ground_Elevation_m
        '',                               # Total_Depth_m
        '',                               # Groundwater_Level_m
        '', '', '', '', '', '', '',       # Topsoil ~ HardRock
        '기본설계 시추주상도(1공구)_45공.pdf' # Source_Report
    ]
    ws.append(row)

# Insert 2공구
for item in borehole_index_2["sections"]:
    row = [
        item['section_id'],               # BH_No
        '기본설계',                        # Stage
        '',                               # Coordinate_X
        '',                               # Coordinate_Y
        item['elevation_m'],              # Ground_Elevation_m
        '',                               # Total_Depth_m
        '',                               # Groundwater_Level_m
        '', '', '', '', '', '', '',       # Topsoil ~ HardRock
        '기본설계 시추주상도(2공구)_27공.pdf' # Source_Report
    ]
    ws.append(row)

wb.save(excel_path)
print(f"Updated {excel_path} with {ws.max_row - 1} borehole rows.")
